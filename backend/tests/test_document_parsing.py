import asyncio
from io import BytesIO

import pytest
from docx import Document
from fastapi import UploadFile
from starlette.datastructures import Headers

from app.exceptions.base import BadRequestException
from app.services import document_service


STANDARD_DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document"
)


def _upload_file(
    content: bytes,
    filename: str,
    content_type: str,
) -> UploadFile:
    return UploadFile(
        filename=filename,
        file=BytesIO(content),
        headers=Headers({"content-type": content_type}),
    )


def _build_docx_content() -> bytes:
    document = Document()
    document.add_heading("Employee Handbook", level=1)
    document.add_paragraph("")
    document.add_paragraph(
        "Employees receive 10 days of paid vacation."
    )

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Project"
    table.rows[0].cells[1].text = "Content"

    vacation_row = table.add_row().cells
    vacation_row[0].text = "Vacation"
    vacation_row[1].text = "10 days"

    sick_leave_row = table.add_row().cells
    sick_leave_row[0].text = "Sick Leave"
    sick_leave_row[1].text = "5 days"

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_parse_markdown_upload_preserves_original_content():
    content = (
        "# Employee Handbook\n\n"
        "## Vacation\n\n"
        "Employees receive 10 days of paid vacation."
    ).encode("utf-8")
    upload = _upload_file(
        content,
        "handbook.md",
        "text/markdown",
    )

    result = asyncio.run(document_service.parse_document(upload))

    assert result == content.decode("utf-8")


def test_parse_markdown_uses_same_utf8_error_behavior_as_txt():
    invalid_content = b"invalid utf-8: \xff"

    with pytest.raises(BadRequestException) as markdown_error:
        asyncio.run(
            document_service.parse_document(
                _upload_file(
                    invalid_content,
                    "handbook.md",
                    "text/markdown",
                )
            )
        )

    with pytest.raises(BadRequestException) as txt_error:
        asyncio.run(
            document_service.parse_document(
                _upload_file(
                    invalid_content,
                    "handbook.txt",
                    "text/plain",
                )
            )
        )

    assert str(markdown_error.value) == str(txt_error.value)


def test_parse_docx_upload_extracts_paragraphs_and_tables():
    upload = _upload_file(
        _build_docx_content(),
        "handbook.docx",
        STANDARD_DOCX_CONTENT_TYPE,
    )

    result = asyncio.run(document_service.parse_document(upload))

    assert result == (
        "Employee Handbook\n"
        "Employees receive 10 days of paid vacation.\n"
        "Project | Content\n"
        "Vacation | 10 days\n"
        "Sick Leave | 5 days"
    )


def test_parse_docx_rejects_invalid_content():
    upload = _upload_file(
        b"not a valid docx",
        "broken.docx",
        STANDARD_DOCX_CONTENT_TYPE,
    )

    with pytest.raises(BadRequestException, match="无法解析DOCX文件"):
        asyncio.run(document_service.parse_document(upload))
